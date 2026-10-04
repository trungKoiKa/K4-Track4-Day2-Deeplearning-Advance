from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd

import eval as ev
from .common import read_json, sha256, write_json
from .dataset import CLASS_NAMES
from .validation import validate_completion


def _table(rows, columns):
    return pd.DataFrame(rows).reindex(columns=columns)


def export_results(root, data_root, student_id=None, student_name=None, notebook_url=''):
    """Export only measured numbers; refuse a submission package until every required run exists."""
    root,data_root=Path(root),Path(data_root)
    lock=read_json(root/'final_lock.json')
    prediction=root/'predictions'; evaluation=root/'eval_out'; evaluation.mkdir(exist_ok=True)
    test_csv=str(data_root/'labels'/'test_subset0.csv'); val_csv=str(data_root/'labels'/'val_subset0.csv')
    labels_csv=str(data_root/'labels'/'labels.csv')
    groups={}
    for exp_id in ('T00','F01'):
        for seed in (0,1,2):
            completion=root/'runs'/exp_id/f'seed{seed}'/'test_complete.json'
            if not completion.exists():
                raise RuntimeError(f'Missing completed test: {exp_id} seed{seed}')
            completed=read_json(completion)
            pred_file=prediction/f'{exp_id}_seed{seed}_test.csv'
            if not pred_file.exists() or completed.get('sha256') != sha256(pred_file):
                raise ValueError(f'Prediction changed after test completion: {exp_id} seed{seed}')
        pattern=str(prediction/f'{exp_id}_seed*_test.csv')
        if ev.main(['score','--pred',pattern,'--test-csv',test_csv,'--labels',labels_csv,'--tag',exp_id,'--out',str(evaluation)])!=0:
            raise ValueError('eval.py score failed')
        group=ev.load_group(pattern,test_csv)
        if sorted(group.seeds)!=[0,1,2]:
            raise ValueError('Expected exactly seeds 0,1,2')
        groups[exp_id]=group
    args=['grade','--final',str(prediction/'F01_seed*_test.csv'),'--baseline',str(prediction/'T00_seed*_test.csv'),
          '--final-val',str(prediction/'F01_seed*_val.csv'),'--val-csv',val_csv,
          '--test-csv',test_csv,'--labels',labels_csv,'--out',str(evaluation)]
    if lock['pipeline']['calibrate']:
        args += ['--uncal',str(prediction/'F01uncal_seed*_test.csv')]
    latency_paths=list((root/'runs'/'F01').glob('seed*/final_latency_batch1.json'))
    if latency_paths:
        latencies=[read_json(p) for p in latency_paths]
        if len({r['gpu'] for r in latencies})==1:
            args += ['--latency-p95-ms',str(max(r['p95'] for r in latencies)),'--latency-method','proper']
    if ev.main(args)!=0:
        raise ValueError('eval.py grade failed')
    backbones=read_json(root/'backbone_comparison.json')
    baseline=read_json(root/'training_baseline.json')
    training=[{'exp_id':'T00','backbone':baseline['backbone'],'axis':'baseline','change':'reused '+baseline['source_exp'],
               'train_seconds_per_epoch':baseline.get('train_seconds_per_epoch'),'seed':0,'macro_f1':baseline['macro_f1'],'top1':baseline['top1'],'delta':0,'notes':'screening, one seed'}]
    for exp_id in ('T01','T02','T03','T04','T05','T06'):
        folder=root/'runs'/exp_id/'seed0'
        result=read_json(folder/'result.json'); details=read_json(folder/'ablation.json')
        training.append({**result,**details,'change':json.dumps(details['change']),'notes':'screening, one seed'})
    for row in training:
        row.update({'p95':baseline['p95'],'latency_source':baseline['source_exp']+'/seed0/latency.json',
                    'latency_note':'Measured shared backbone single-view FP32; linked, not remeasured per recipe'})
    inference=read_json(root/'inference_comparison.json')
    final=[]; perclass=[]; latency=[]
    for exp_id,group in groups.items():
        val_scores=[]; val_top1=[]
        for pred,metrics in zip(group.preds,group.metrics):
            folder=root/'runs'/exp_id/f'seed{pred.seed}'
            val_pred=ev.read_pred(str(prediction/f'{exp_id}_seed{pred.seed}_val.csv'))
            ev.check_against_csv(val_pred,val_csv,'val')
            val=ev.compute_metrics(val_pred.y_true,val_pred.y_pred,val_pred.probs)
            val_scores.append(val['macro_f1']); val_top1.append(val['top1'])
            cfg=read_json(folder/'config.json')
            final.append({'exp_id':exp_id,'seed':pred.seed,'config':json.dumps({'backbone':cfg['backbone'],'init':cfg['init'],'mix':cfg['mix'],'loss':cfg['loss'],'pipeline':lock['pipeline'] if exp_id=='F01' else 'single FP32'}),
                          'macro_f1_val':val['macro_f1'],'top1_val':val['top1'],'macro_f1_test':metrics['macro_f1'],'top1_test':metrics['top1'],
                          'ece_test':metrics['ece'],'balanced_acc_test':metrics['balanced_acc']})
        final.append({'exp_id':exp_id,'seed':'mean ± sample std','config':'3 seeds',
                      'top1_val':float(np.mean(val_top1)),'top1_val_std':float(np.std(val_top1,ddof=1)),
                      'macro_f1_val':float(np.mean(val_scores)), 'macro_f1_val_std':float(np.std(val_scores,ddof=1)),
                      **{key:group.summary[source][0] for key,source in [('macro_f1_test','macro_f1'),('top1_test','top1'),('ece_test','ece'),('balanced_acc_test','balanced_acc')]},
                      **{key+'_std':group.summary[source][1] for key,source in [('macro_f1_test','macro_f1'),('top1_test','top1'),('ece_test','ece'),('balanced_acc_test','balanced_acc')]}})
        for label,name in enumerate(CLASS_NAMES):
            perclass.append({'exp_id':exp_id,'class':name,'test_images':int(group.metrics[0]['support'][label]),
                             **{key:float(group.summary[key][0][label]) for key in ('precision','recall','f1')},
                             **{key+'_std':float(group.summary[key][1][label]) for key in ('precision','recall','f1')}})
    for path in list((root/'inference').glob('*_latency.json'))+list((root/'runs').glob('*/seed*/final_latency_batch*.json')):
        row=read_json(path); row={k:v for k,v in row.items() if k!='samples_ms'}
        row.update({'config':str(path.relative_to(root)),'source_log':str(path.relative_to(root))})
        latency.append(row)
    summary=[{'exp_id':r['exp_id'],'phase':'backbone','n_seeds':1,'train_seconds_per_epoch':r.get('train_seconds_per_epoch'),'macro_f1_val':r['macro_f1'],'p95_ms':r['p95'],'notes':'one screening seed'} for r in backbones]
    summary += [{'exp_id':r['exp_id'],'phase':'training','n_seeds':1,'train_seconds_per_epoch':r.get('train_seconds_per_epoch'),'p95_ms':r['p95'],'cost_source':r['latency_source'],'macro_f1_val':r['macro_f1'],'notes':r['change']} for r in training]
    summary += [{'exp_id':r['exp_id'],'phase':'inference','n_seeds':1,'macro_f1_val':r['macro_f1'],'p95_ms':r['p95'],'relative_cost':r['relative_cost']} for r in inference]
    comparison=[r for r in final if r['seed']=='mean ± sample std']
    for row in comparison:
        paths=list((root/'runs'/row['exp_id']).glob('seed*/final_latency_batch1.json'))
        timings=[read_json(p) for p in paths]
        same_gpu=timings and len({r['gpu'] for r in timings})==1
        summary.append({'exp_id':row['exp_id'],'phase':'final','n_seeds':3,'macro_f1_val':row['macro_f1_val'],
                        'p95_ms':max(r['p95'] for r in timings) if same_gpu else None,
                        'cost_source':'final measured latency logs','notes':'mean val across 3 seeds; hardware groups in Latency'})
    summary=sorted(summary,key=lambda r:-r['macro_f1_val'])[:10]
    sheets={'Summary':_table(summary,['exp_id','phase','n_seeds','macro_f1_val','p95_ms','relative_cost','train_seconds_per_epoch','cost_source','notes']),
            'Backbones':pd.DataFrame([{**{k:v for k,v in r.items() if k!='config'},'seed':0,
                         'notes':'MAC exclusions: '+json.dumps(read_json(root/'runs'/r['exp_id']/'seed0'/'profile.json')['unsupported_ops'])} for r in backbones]),
            'Training':pd.DataFrame(training),'Inference':pd.DataFrame(inference),
            'Final':pd.DataFrame(final),'PerClass':pd.DataFrame(perclass),'Latency':pd.DataFrame(latency)}
    for sheet in ('Backbones','Training','Inference'):
        sheets[sheet]=sheets[sheet].rename(columns={'macro_f1':'macro_f1_val','top1':'top1_val',
                                                  'ece':'ece_val','p50':'p50_ms','p95':'p95_ms','p99':'p99_ms'})
    sheets['Latency']=sheets['Latency'].rename(columns={'p50':'p50_ms','p95':'p95_ms','p99':'p99_ms'})
    # The requested implementation explicitly uses pandas/openpyxl on Colab.
    with pd.ExcelWriter(root/'results.xlsx',engine='openpyxl') as writer:
        for name,df in sheets.items():
            df.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name]; ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Font, PatternFill, Alignment
            for cell in ws[1]:
                cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='17365D')
                cell.alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height=35
            for column in ws.columns:
                width=max(len(str(c.value or '')) for c in list(column)[:30])
                ws.column_dimensions[column[0].column_letter].width=min(45,max(14,width+2))
                for cell in list(column)[1:]:
                    if isinstance(cell.value,float): cell.number_format='0.0000'
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
            if 'macro_f1' in df or 'macro_f1_val' in df:
                metric='macro_f1' if 'macro_f1' in df else 'macro_f1_val'
                if df[metric].notna().any():
                    row=int(df[metric].idxmax())+2
                    for cell in ws[row]: cell.fill=PatternFill('solid',fgColor='E2F0D9')
    from openpyxl import load_workbook
    wb=load_workbook(root/'results.xlsx'); ws=wb['Summary']
    start=ws.max_row+3; ws.cell(start,1,'Baseline vs final (3 seeds; mean and sample std)')
    cols=['exp_id','macro_f1_val','macro_f1_val_std','top1_val','top1_val_std','macro_f1_test','macro_f1_test_std','top1_test','top1_test_std','ece_test','ece_test_std']
    ws.append(cols)
    for row in comparison: ws.append([row.get(k) for k in cols])
    wb.save(root/'results.xlsx')
    _backbone_figures(root,backbones)
    _analysis_figures(root,data_root,groups)
    _report(root,groups,backbones,training,inference,lock,notebook_url)
    _readme(root,notebook_url)
    write_json(root/'export_provenance.json',{'prediction_sha256':{p.name:sha256(p) for p in prediction.glob('*_test.csv')},
               'workbook_sha256':sha256(root/'results.xlsx'),'eval_sha256':sha256(Path(ev.__file__))})
    status={**validate_completion(root,data_root,notebook_url),'results_exported':True,'human_analysis_complete':(root/'analysis_notes.md').exists() and len((root/'analysis_notes.md').read_text(encoding='utf-8').strip())>=200,
            'required_files':['README.md','results.xlsx','report.md','curves','predictions','code']}
    write_json(root/'submission_status.json',status)
    if student_id or student_name:
        if not student_id or not student_name or not re.fullmatch(r'[A-Za-z0-9_-]+',student_id) or not re.fullmatch(r'[A-Za-z0-9_]+',student_name):
            raise ValueError('MSSV and unaccented underscore-separated name required')
        if not status['human_analysis_complete']:
            raise ValueError('Review error images and write >=200 characters of analysis_notes.md before packaging')
        if not status['ready_to_submit']:
            raise ValueError('Submission incomplete: '+ '; '.join(status['errors']))
        destination=root/'submissions'/f'{student_id}_{student_name}'
        destination.mkdir(parents=True,exist_ok=True)
        for name in ('README.md','results.xlsx','report.md','analysis_notes.md','final_lock.json','export_provenance.json','runtime_versions.json','environment.lock.txt'):
            if not (root/name).exists():
                continue
            shutil.copy2(root/name,destination/name)
        for path in root.glob('*.json'):
            shutil.copy2(path,destination/path.name)
        for name in ('curves','predictions','eval_out','eda','pipeline','analysis','inference'):
            if (root/name).exists(): shutil.copytree(root/name,destination/name,dirs_exist_ok=True)
        code=destination/'code'; code.mkdir(exist_ok=True)
        source=Path(__file__).parent
        shutil.copytree(source,code/'lab_solution',ignore=shutil.ignore_patterns('__pycache__'),dirs_exist_ok=True)
        shutil.copy2(ev.__file__,code/'eval.py')
        for name in ('requirements-colab.txt','README_IMPLEMENTATION.md','01_DeepWeeds_Full_Lab_Colab.ipynb'):
            if (source.parent/name).exists(): shutil.copy2(source.parent/name,code/name)
        # Include evidence needed to trace every training number; never include weights or image data.
        for source_file in (root/'runs').rglob('*'):
            if source_file.is_file() and source_file.suffix in ('.json','.csv'):
                target=destination/'logs'/source_file.relative_to(root/'runs')
                target.parent.mkdir(parents=True,exist_ok=True); shutil.copy2(source_file,target)
        shutil.copytree(source.parent/'solution_tests',code/'solution_tests',ignore=shutil.ignore_patterns('__pycache__'),dirs_exist_ok=True)
        zip_path=shutil.make_archive(str(destination),'zip',root_dir=destination)
        return {'submission':str(destination),'zip':zip_path,'status':status}
    return {'results':str(root/'results.xlsx'),'report':str(root/'report.md'),'status':status}


def _final_val_table(root):
    rows=['| Config | Macro-F1 val | Top-1 val |','| --- | --- | --- |']
    for exp in ('T00','F01'):
        metrics=[]
        for seed in range(3):
            p=ev.read_pred(str(root/'predictions'/f'{exp}_seed{seed}_val.csv'))
            metrics.append(ev.compute_metrics(p.y_true,p.y_pred,p.probs))
        values=[]
        for key in ('macro_f1','top1'):
            x=[m[key] for m in metrics]
            values.append(f'{np.mean(x):.4f} ± {np.std(x,ddof=1):.4f}')
        rows.append('| '+exp+' | '+' | '.join(values)+' |')
    return '\n'.join(rows)


def _backbone_figures(root,rows):
    import matplotlib.pyplot as plt
    folder=root/'analysis'; folder.mkdir(exist_ok=True)
    for key,label,name in [('p95','Batch-1 p95 (ms)','backbone_latency'),('params_m','Parameters (M)','backbone_params')]:
        fig,ax=plt.subplots(figsize=(7,5))
        for row in rows:
            ax.scatter(row[key],row['macro_f1']); ax.annotate(row['exp_id'],(row[key],row['macro_f1']))
        ax.set_xlabel(label); ax.set_ylabel('Val macro-F1 (screening seed0)'); ax.grid(alpha=.2)
        fig.tight_layout(); fig.savefig(folder/(name+'.png'),dpi=160); plt.close(fig)


def _analysis_figures(root,data_root,groups):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from PIL import Image
    folder=root/'analysis'; folder.mkdir(exist_ok=True)
    for exp_id,group in groups.items():
        cm=sum(m['confusion'] for m in group.metrics)
        fig,ax=plt.subplots(figsize=(10,8)); im=ax.imshow(cm,cmap='Blues'); fig.colorbar(im,ax=ax)
        ax.set_xticks(range(9),CLASS_NAMES,rotation=90); ax.set_yticks(range(9),CLASS_NAMES)
        for i in range(9):
            for j in range(9): ax.text(j,i,str(cm[i,j]),ha='center',va='center',fontsize=8,color='white' if cm[i,j]>cm.max()/2 else 'black')
        ax.set_xlabel('Predicted'); ax.set_ylabel('True'); ax.set_title(exp_id+' confusion counts summed over 3 seeds')
        fig.tight_layout(); fig.savefig(folder/f'{exp_id}_confusion.png',dpi=160); plt.close(fig)
    # Prespecified seed0 error review; never select the best test seed.
    pred=next(p for p in groups['F01'].preds if p.seed==0)
    wrong=np.flatnonzero(pred.y_true!=pred.y_pred)
    hard=[i for i in wrong if pred.y_true[i] in (0,7) and pred.y_pred[i] in (0,7)]
    other=[i for i in wrong if i not in hard]
    chosen=(hard+other)[:12]
    records=[{'Filename':str(pred.filenames[i]),'y_true':int(pred.y_true[i]),'y_pred':int(pred.y_pred[i]),
              'confidence':float(pred.probs[i].max()),'observation':'','hypothesis':''} for i in chosen]
    pd.DataFrame(records).to_csv(folder/'error_review_seed0.csv',index=False)
    if chosen:
        fig,axes=plt.subplots(3,4,figsize=(15,11))
        for ax,i in zip(axes.flat,chosen):
            with Image.open(data_root/'images'/str(pred.filenames[i])) as image: ax.imshow(image.convert('RGB'))
            ax.set_title(f'True: {CLASS_NAMES[pred.y_true[i]]}\nPred: {CLASS_NAMES[pred.y_pred[i]]}',fontsize=9); ax.axis('off')
        for ax in list(axes.flat)[len(chosen):]: ax.axis('off')
        fig.tight_layout(); fig.savefig(folder/'error_images_seed0.png',dpi=140); plt.close(fig)


def _report(root,groups,backbones,training,inference,lock,notebook_url):
    f,b=groups['F01'],groups['T00']
    def metric(group,key):
        mean,std=group.summary[key]
        return f'{mean:.4f} ± {std:.4f}'
    delta=f.summary['macro_f1'][0]-b.summary['macro_f1'][0]
    noise=max(f.summary['macro_f1'][1],b.summary['macro_f1'][1])
    conclusion='vượt độ lệch chuẩn lớn hơn của hai nhóm' if delta>noise else 'chưa vượt nhiễu giữa seed; không khẳng định tốt hơn'
    def table(rows,cols):
        df=pd.DataFrame(rows).reindex(columns=cols)
        lines=['| '+' | '.join(cols)+' |','| '+' | '.join(['---']*len(cols))+' |']
        for _,row in df.iterrows():
            lines.append('| '+' | '.join(f'{v:.4f}' if isinstance(v,float) else str(v).replace('|','/') for v in row)+' |')
        return '\n'.join(lines)
    notes=(root/'analysis_notes.md').read_text(encoding='utf-8') if (root/'analysis_notes.md').exists() else '**Cần bổ sung nhận xét quan sát ảnh mẫu, ảnh lỗi và giả thuyết nguyên nhân vào analysis_notes.md.**'
    architecture_delta=max(r['macro_f1'] for r in backbones)-next(r['macro_f1'] for r in backbones if r['exp_id']=='B01')
    training_delta=max(r['macro_f1'] for r in training)-training[0]['macro_f1']
    inference_delta=max(r['macro_f1'] for r in inference)-inference[0]['macro_f1']
    contributions={'backbone vs ResNet50':architecture_delta,'recipe vs T00':training_delta,'inference vs I00':inference_delta}
    largest=max(contributions,key=contributions.get)
    combination=next(r for r in training if r['exp_id']=='T06')
    combination_change=json.loads(combination['change'])
    loss_id='T04' if combination_change['loss']=='ls' else 'T05'
    additive=next(r['delta'] for r in training if r['exp_id']=='T03')+next(r['delta'] for r in training if r['exp_id']==loss_id)
    interaction=combination['delta']-additive
    curves=[]
    for path in sorted((root/'curves').glob('*.png')):
        history_path=root/'runs'/path.stem.split('_seed')[0]/('seed'+path.stem.rsplit('_seed',1)[1])/'history.csv'
        if history_path.exists():
            history=pd.read_csv(history_path)
            best_epoch=int(history.loc[history.macro_f1.idxmax(),'epoch'])
            fall=float(history.macro_f1.max()-history.macro_f1.iloc[-1])
            curves.append(f'- {path.stem}: checkpoint epoch{best_epoch}, macro-F1 cuối giảm {fall:.4f} so với đỉnh; xem ![curve]({path.relative_to(root).as_posix()}).')
    final_timings=[read_json(p) for p in (root/'runs'/'F01').glob('seed*/final_latency_batch1.json')]
    realtime='Chưa có latency GPU của cấu hình cuối; cần đo trước kết luận triển khai.'
    if final_timings:
        measured=max(r['p95'] for r in final_timings)
        realtime=f'p95 lớn nhất thực đo của F01 qua các phiên = {measured:.2f}ms: '+('đạt' if measured<=100 else 'không đạt')+' mốc100ms trong phạm vi GPU inference. Phần cứng: '+', '.join(sorted({r['gpu'] for r in final_timings}))+'. Không gộp timing từ GPU khác thành một phân phối.'
    text=f'''# Báo cáo Lab Day 2 — DeepWeeds

## 1. Tóm tắt
So sánh 5 backbone, 3 trục huấn luyện và 4 phương pháp suy luận ngoài mốc.
Cấu hình chung kết: `{lock['final_config']['backbone']}`, recipe được chọn trên val, pipeline `{lock['pipeline']}`.
Test macro-F1: **{metric(f,'macro_f1')}**, top-1: **{metric(f,'top1')}**, qua 3 seed (std mẫu ddof=1).
Baseline macro-F1: **{metric(b,'macro_f1')}**. Δ={delta:.4f}, {conclusion}.
Mọi số dưới đây được tính từ dự đoán thật bằng eval.py gốc.

## 2. Dữ liệu và thiết lập
DeepWeeds, 17.509 ảnh, 9 lớp; fold 0 nguyên bản, train 10.501 / val 3.501 / test 3.507.
Kiểm tra giao rỗng, hợp đầy đủ, tồn tại ảnh và phân bố lớp nằm trong eda/split_checks.json.
Negative có 9.106 ảnh (~52%); macro-F1 là chỉ số chính để tránh lớp đông chi phối accuracy.
Một nhãn không nhất quán có sẵn trong nguồn: 20170714-110407-3.jpg, train nhãn 0 và labels.csv nhãn 1. Giữ nguyên CSV, Dataset dùng nhãn từng split.
Train crop/lật ngang, val resize 256 và center crop 224; chuẩn hóa ImageNet.
AdamW, LR backbone 1e-4/head 1e-3, decay 0.05 trừ norm/bias, warmup 1 epoch và cosine, AMP, 10 epoch.
Batch và accumulation, tag pretrained, phiên bản và phần cứng: xem config.json/environment.json của từng run.
Chọn checkpoint macro-F1 val lớn nhất; hòa chọn epoch sớm. Test chỉ dùng sau final_lock.json, một lượt hoàn tất mỗi cấu hình/seed, có checkpoint từng batch để tiếp tục khi gián đoạn.

![Phân bố lớp](eda/class_distribution.png)
![Mẫu train](eda/train_samples_3_per_class.png)
![Overfit kiểm tra](pipeline/overfit_batch.png)
![Augmentation và CutMix](pipeline/augmentation_cutmix.png)
Bằng chứng định lượng: {json.dumps(read_json(root/'pipeline_checks.json'),ensure_ascii=False) if (root/'pipeline_checks.json').exists() else 'Thiếu pipeline checks'}.
Nhật ký thất bại: {json.dumps(read_json(root/'failures.json'),ensure_ascii=False) if (root/'failures.json').exists() else 'Không có thất bại được ghi nhận' }.

## 3. So sánh backbone
{table(backbones,['exp_id','backbone','params_m','gmac','macro_f1','top1','train_seconds_per_epoch','p95'])}
![Backbone F1–latency](analysis/backbone_latency.png)
![Backbone F1–params](analysis/backbone_params.png)
Sàng lọc dùng một seed, không khẳng định ý nghĩa thống kê của chênh lệch nhỏ.
Các timing để xếp hạng được đo trên cùng GPU; GMAC dùng fvcore với danh sách toán tử chưa đếm ghi trong log, không coi GMAC là độ trễ.

## 4. Công thức huấn luyện
{table(training,['exp_id','axis','change','macro_f1','delta'])}
Mỗi ablation đổi một yếu tố so với T00; T06 là kết hợp CutMix và loss, có ghi nguồn lựa chọn.
T00 sàng lọc tái sử dụng lần chạy backbone được chọn; baseline chung kết train lại riêng với 3 seed.
T06 dùng loss {loss_id}; Δ kết hợp = {combination['delta']:.4f}, tổng Δ từng yếu tố = {additive:.4f}, phần tương tác = {interaction:.4f}. Đây là quan sát một seed, chưa phải kiểm định cộng dồn.
Đường cong loss train là loss recipe, val loss luôn CE; không so sánh trực tiếp hai loại loss như cùng thước đo.
Kết luận nhiều seed chỉ áp dụng cho chung kết và baseline, không gán std của chung kết cho từng ablation.

### Đường cong và hội tụ
{chr(10).join(curves)}

## 5. Suy luận và hiệu chuẩn
{table(inference,['exp_id','method','K','amp','macro_f1','top1','ece','p50','p95','p99','relative_cost'])}
Five-crop lấy từ ảnh resize 256 trước center crop; TTA gộp logit.
Nhiệt độ chỉ khớp trên val; nếu được bật bởi val seed0, từng seed chung kết khớp T riêng theo cùng chính sách đã khóa.
ECE val trước TS (I00)={inference[0]['ece']:.4f}, sau TS (I04)={inference[4]['ece']:.4f}.
ECE 15 bin. Timing: 10 warmup, 100 lần đo, đồng bộ GPU, batch 1/32; gồm forward, gộp view và softmax/hiệu chuẩn, không gồm đọc ảnh/CPU preprocessing.

![Đánh đổi chất lượng và độ trễ](inference/accuracy_latency.png)

## 6. Chung kết và phân tích lỗi
| Cấu hình | Macro-F1 test | Top-1 test | Balanced acc test | ECE test |
| --- | --- | --- | --- | --- |
| T00 | {metric(b,'macro_f1')} | {metric(b,'top1')} | {metric(b,'balanced_acc')} | {metric(b,'ece')} |
| F01 | {metric(f,'macro_f1')} | {metric(f,'top1')} | {metric(f,'balanced_acc')} | {metric(f,'ece')} |

Val qua 3 seed (mean ± sample std):
{_final_val_table(root)}

Recall Chinee apple: {f.summary['recall'][0][0]:.4f} ± {f.summary['recall'][1][0]:.4f}; Snake weed: {f.summary['recall'][0][7]:.4f} ± {f.summary['recall'][1][7]:.4f}.
PerClass trong Excel và eval_out chứa precision/recall/F1 của đủ 9 lớp, mean ± std.
Ma trận dưới là số đếm cộng qua 3 seed (3 × 3.507 lượt ảnh); hàng nhãn thật, cột dự đoán.

![Ma trận nhầm lẫn](analysis/F01_confusion.png)
![Ảnh lỗi seed0](analysis/error_images_seed0.png)

### Nhận xét và giả thuyết từ quan sát
{notes}

## 7. Kết luận và khuyến nghị
Δ macro-F1 chung kết–baseline = {delta:.4f}; std lớn hơn = {noise:.4f}: {conclusion}.
Đối chiếu mức thay đổi của backbone, recipe và inference bằng các bảng trên; các lần sàng lọc một seed chỉ mang tính gợi ý.
Mức thay đổi sàng lọc: kiến trúc so với B01={architecture_delta:.4f}; recipe so với T00={training_delta:.4f}; inference so với I00={inference_delta:.4f}. Yếu tố có Δ lớn nhất trong thiết kế này: {largest}. Các mốc so sánh khác nhau và trọng số pretrained khác recipe nên không coi đây là phân rã nhân quả độc lập.
Khuyến nghị robot dựa trên p95 batch 1 của pipeline thực đo; mốc 100 ms chỉ nói về phạm vi GPU inference đã khai báo, không bảo đảm độ trễ toàn hệ thống cảm biến.
{realtime}
Nếu p95 vượt 100 ms, báo không đạt ngân sách, không quay lại chọn bằng test.

## 8. Hạn chế
Colab miễn phí, 10 epoch, 1 fold, 3 seed chung kết; nhiều thí nghiệm sàng lọc chỉ 1 seed.
Chia ngẫu nhiên không theo địa điểm có thể làm test lạc quan khi triển khai ở địa điểm/mùa/ánh sáng mới.
Gradient accumulation không tương đương batch 64 đối với thống kê BatchNorm; tất cả backbone giữ cùng batch thực để so sánh công bằng.
GPU có thể thay đổi giữa phiên; timing chỉ so sánh trực tiếp trong cùng nhóm phần cứng.
ConvNeXt dùng fb_in1k; tất cả backbone dùng ImageNet-1k nhưng công thức pretrained khác nhau. Xem tag đầy đủ trong config/log.
Số accuracy bài báo (95.7% ResNet-50, 95.1% Inception-v3) là trích dẫn tham khảo từ README, không phải kết quả của bài này; điều kiện train khoảng 100 epoch khác lab.
Không làm điểm thưởng và không hứa đạt mốc accuracy của bài báo.

## 9. Phụ lục và tái lập
Notebook: {notebook_url or 'Điền link notebook Colab của bạn trong README.'}
Code/cấu hình/log/curves/predictions liên kết bằng exp_id và seed; hash dữ liệu/code/checkpoint được ghi để kiểm tra nguồn số liệu.
eval_out chứa kết quả score/grade; final_lock.json ghi quyết định trước test. Không đưa dataset/checkpoint lớn vào Git.
'''
    (root/'report.md').write_text(text,encoding='utf-8')


def _readme(root,notebook_url):
    text=f'''# Bài làm DeepWeeds — hướng dẫn chạy lại

Notebook Colab: {notebook_url or 'Điền link notebook của bạn'}

1. Chạy notebook EDA hoặc phần chuẩn bị trong notebook điều khiển.
2. Cài requirements-colab.txt, giữ PyTorch/CUDA của Colab; lưu environment.json và environment.lock.txt sau kiểm tra.
3. Chạy EDA → pipeline_checks → preflight → B01..B05 → select_backbone.
4. Chạy T01..T06 → select_recipe → compare_inference → lock_final.
5. Train T00/F01 mỗi seed 0,1,2; test mỗi cấu hình/seed một lượt hoàn tất.
6. export tạo Excel/báo cáo; đọc ảnh lỗi và thêm analysis_notes.md rồi export lại để đóng gói.

Khi phiên ngắt: kết nối Drive, chuẩn bị lại ảnh vào /content, cài đúng phiên bản và tạo Study với cùng root. Chạy lại ô thí nghiệm; resume checkpoint gần nhất. Dự đoán test đã lưu được tái sử dụng, không rerun để chọn điểm.
Seed sàng lọc: 0. Seed chung kết/baseline: 0,1,2. Fold: 0, không sửa CSV/gộp val.
Đường dẫn mặc định Drive: Lab_Day2_DeepWeeds/study. Dữ liệu: /content/deepweeds.
Code chạy qua một hàm train.run(Config); không sửa eval.py. Phiên bản, tag, hyperparameter và phần cứng có trong logs của mỗi run.
Checkpoint lớn nằm trên Drive, không có trong ZIP bài nộp. logs chứa history/config/environment để truy số liệu.
'''
    (root/'README.md').write_text(text,encoding='utf-8')
