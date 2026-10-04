"""Export contract test. All numeric fixtures are synthetic and confined to a temporary directory."""
import json
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image
from openpyxl import load_workbook

import eval as ev
from lab_solution.common import write_json, sha256
from lab_solution.reporting import export_results


class ExportTests(unittest.TestCase):
    def test_export_recomputes_workbook_from_predictions(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)/'study'; data=Path(temp)/'data'
            root.mkdir(); (data/'labels').mkdir(parents=True); (data/'images').mkdir()
            names=[f'{i}.jpg' for i in range(9)]; y=np.arange(9)
            df=pd.DataFrame({'Filename':names,'Label':y,'Species':ev.CLASS_NAMES})
            for filename in ('labels.csv','test_subset0.csv','val_subset0.csv'): df.to_csv(data/'labels'/filename,index=False)
            for name in names: Image.new('RGB',(20,20)).save(data/'images'/name)
            backbone={'backbone':'resnet50','macro_f1':.5,'top1':.5,'params_m':1.,'gmac':.1,'train_seconds_per_epoch':1.,'p95':5.,'p50':4.}
            rows=[]
            for i in range(1,6):
                exp=f'B{i:02d}'; rows.append({'exp_id':exp,**backbone})
                write_json(root/'runs'/exp/'seed0'/'profile.json',{'unsupported_ops':{}})
            write_json(root/'backbone_comparison.json',rows)
            write_json(root/'training_baseline.json',{'source_exp':'B01',**backbone})
            for i in range(1,7):
                exp=f'T{i:02d}'; folder=root/'runs'/exp/'seed0'
                write_json(folder/'result.json',{'exp_id':exp,'seed':0,**backbone})
                write_json(folder/'ablation.json',{'axis':'test fixture','change':{'mix':'cutmix','loss':'ls'} if i==6 else {'loss':'ce'},'delta':0.})
            inference=[]
            for i in range(5):
                inference.append({'exp_id':f'I{i:02d}','method':'single','K':1,'amp':False,'macro_f1':.5,'top1':.5,'ece':.2,
                                  'p50':4.,'p95':5.,'p99':6.,'relative_cost':1.})
            write_json(root/'inference_comparison.json',inference)
            write_json(root/'final_lock.json',{'final_config':{'backbone':'resnet50'},'pipeline':{'method':'single','amp':False,'calibrate':False}})
            for exp in ('T00','F01'):
                for seed in range(3):
                    p=np.eye(9)*.8+np.ones((9,9))*.2/9
                    if exp=='T00': p=np.roll(p,seed,axis=1)
                    folder=root/'runs'/exp/f'seed{seed}'
                    write_json(folder/'final_calibration.json',{'val_metrics':{'macro_f1':.5}})
                    write_json(folder/'config.json',{'backbone':'resnet50','init':'finetune','mix':None,'loss':'ce'})
                    for split in ('test','val'): ev.save_predictions(root/'predictions'/f'{exp}_seed{seed}_{split}.csv',names,y,p)
                    write_json(folder/'test_complete.json',{'sha256':sha256(root/'predictions'/f'{exp}_seed{seed}_test.csv')})
            result=export_results(root,data,notebook_url='https://example.test/synthetic-fixture')
            self.assertFalse(result['status']['human_analysis_complete'])
            self.assertFalse(result['status']['ready_to_submit'])
            self.assertTrue(any('curves/' in e for e in result['status']['errors']))
            self.assertTrue(any('final_latency' in e for e in result['status']['errors']))
            wb=load_workbook(root/'results.xlsx',data_only=True)
            self.assertEqual(set(wb.sheetnames),{'Summary','Backbones','Training','Inference','Final','PerClass','Latency'})
            ws=wb['Final']; headers=[c.value for c in ws[1]]
            values=[dict(zip(headers,row)) for row in ws.iter_rows(min_row=2,values_only=True)]
            total=next(r for r in values if r['exp_id']=='F01' and r['seed']=='mean ± sample std')
            group=ev.load_group(str(root/'predictions'/'F01_seed*_test.csv'),str(data/'labels'/'test_subset0.csv'))
            self.assertAlmostEqual(total['macro_f1_test'],group.summary['macro_f1'][0])
            self.assertAlmostEqual(total['macro_f1_test_std'],group.summary['macro_f1'][1])
            self.assertAlmostEqual(total['top1_val'],1.)
            self.assertAlmostEqual(total['top1_val_std'],0.)
            summary_values=[c.value for row in wb['Summary'] for c in row]
            self.assertIn('Baseline vs final (3 seeds; mean and sample std)',summary_values)
            self.assertTrue((root/'analysis'/'backbone_latency.png').exists())
            self.assertTrue((root/'analysis'/'backbone_params.png').exists())
            self.assertEqual(ws.freeze_panes,'A2')
            self.assertIn('có sẵn trong nguồn',(root/'report.md').read_text(encoding='utf-8'))
            with self.assertRaisesRegex(ValueError,'analysis_notes'): export_results(root,data,'123','Nguyen_Van_A')
            (root/'analysis_notes.md').write_text('Synthetic test observations. '*20,encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'Submission incomplete'): export_results(root,data,'123','Nguyen_Van_A')
            modified=root/'predictions'/'F01_seed0_test.csv'
            modified.write_text(modified.read_text()+'\n',encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'changed after test completion'): export_results(root,data)


if __name__=='__main__': unittest.main()
